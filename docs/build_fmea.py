"""
build_fmea.py -- Generates FMEA_RiskAudit.xlsx for Review 3 (Step 5).
FMEA = Failure Mode & Effects Analysis. RPN = Severity x Occurrence x Detection (formula-driven).
Scope: Hospital Management System, Quality Goal Q07 (Improve Data Accuracy).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ---------------------------------------------------------------- FMEA sheet
ws = wb.active
ws.title = "FMEA"

headers = ["Process Step / Feature", "Potential Failure Mode", "Potential Effect(s)",
           "Severity (1-10)", "Potential Cause(s)", "Occurrence (1-10)",
           "Current Controls", "Detection (1-10)", "RPN", "Recommended Action"]

rows = [
    ["Patient Phone Entry", "Digits mistyped / wrong count", "Unreachable patient, missed follow-up",
     7, "Manual free-text entry", 6, "None (pre-fix)", 8, None,
     "Add Input Masking (+91-XXXXX-XXXXX) with 10-digit enforcement"],
    ["Gender / Blood Group Entry", "Free-text typo (e.g. 'mael', 'O+ ve')", "Incorrect medical record, transfusion risk",
     9, "Free-text field, no constraint", 5, "None (pre-fix)", 6, None,
     "Replace with fixed Dropdown List (Combobox, read-only)"],
    ["New Patient Registration", "Duplicate patient record created", "Fragmented history, billing errors",
     6, "No search-before-add step", 6, "None (pre-fix)", 7, None,
     "Add Auto-Complete search to surface existing matches before insert"],
    ["Patient/Appointment Delete", "Accidental deletion of correct record", "Irrecoverable data loss",
     8, "Single-click delete, no confirmation", 4, "None (pre-fix)", 3, None,
     "Add Confirmation Modal (Yes/No) before every delete/critical update"],
    ["Any Data Modification", "Unauthorized or unexplained record change", "No traceability for audits",
     7, "No change history kept", 5, "None (pre-fix)", 9, None,
     "Add Audit Log capturing table, record, action, user, timestamp"],
    ["Date of Birth Entry", "Invalid / impossible date format", "Age-based clinical logic fails",
     6, "Unvalidated text field", 5, "Regex validation (post-fix)", 4, None,
     "Enforce YYYY-MM-DD regex validation on submit"],
]

bold = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4472C4")
thin = Side(style="thin", color="AAAAAA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = bold
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

for r in rows:
    ws.append(r)

# RPN formula = Severity * Occurrence * Detection  (columns D, F, H -> I)
last_row = ws.max_row
for r in range(2, last_row + 1):
    ws.cell(row=r, column=9).value = f"=D{r}*F{r}*H{r}"

widths = [22, 26, 28, 12, 26, 12, 22, 12, 8, 34]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for r in range(2, last_row + 1):
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)

ws.freeze_panes = "A2"

# Conditional-style manual highlight note (High RPN needing priority action)
note_row = last_row + 2
ws.cell(row=note_row, column=1,
        value="Note: Rows with RPN >= 250 are highest priority per FMEA convention "
              "(Severity x Occurrence x Detection, max 1000). Re-score Detection after each fix "
              "and re-run this sheet each PDCA cycle.").font = Font(italic=True, size=9)

# ---------------------------------------------------------- Checksheet sheet
ws2 = wb.create_sheet("Defect Checksheet")
cs_headers = ["Date", "Category", "Description", "Severity", "Status"]
ws2.append(cs_headers)
for c in range(1, len(cs_headers) + 1):
    cell = ws2.cell(row=1, column=c)
    cell.font = bold
    cell.fill = header_fill
    cell.border = border

sample_defects = [
    ["2026-08-05", "Data Entry", "Patient phone saved with 9 digits", "High", "Fixed"],
    ["2026-08-06", "Data Entry", "Gender typed as 'mael'", "High", "Fixed"],
    ["2026-08-07", "Validation", "DOB accepted in DD/MM/YYYY causing wrong age calc", "Medium", "Fixed"],
    ["2026-08-08", "Data Entry", "Duplicate patient 'Ramesh Singh' created twice", "High", "Fixed"],
    ["2026-08-09", "UI", "Delete button triggered with no confirmation", "Critical", "Fixed"],
    ["2026-08-10", "Data Entry", "Blood group left blank, saved as NULL", "High", "Fixed"],
    ["2026-08-11", "Performance", "Patient list slow to load with 500+ rows", "Low", "Open"],
    ["2026-08-12", "Validation", "Appointment time accepted as '25:99'", "Medium", "Fixed"],
    ["2026-08-13", "UI", "Dropdown showed blank first option", "Low", "Fixed"],
    ["2026-08-14", "Data Entry", "Phone number pasted with country code twice", "Medium", "Fixed"],
    ["2026-08-15", "Other", "App title showed wrong quality goal code", "Low", "Fixed"],
]
for row in sample_defects:
    ws2.append(row)
for i, w in enumerate([12, 14, 42, 12, 10], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for r in range(1, ws2.max_row + 1):
    for c in range(1, len(cs_headers) + 1):
        ws2.cell(row=r, column=c).border = border

import os

out_dir = os.path.dirname(os.path.abspath(__file__))
out_path = os.path.join(out_dir, "FMEA_RiskAudit.xlsx")
wb.save(out_path)
print(f"Saved {out_path}")
